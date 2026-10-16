"""
文档解析模块
支持多种格式文档的解析、清洗和分块
"""
import os
import logging
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime
import re

logger = logging.getLogger(__name__)

class DocumentParser:
    """文档解析器"""
    
    def __init__(self):
        self.supported_extensions = {'.pdf', '.txt', '.doc', '.docx', '.xls', '.xlsx', '.md'}
        self.max_file_size = 10 * 1024 * 1024  # 10MB
    
    def parse_file(self, file_path: str) -> Dict[str, Any]:
        """解析文件"""
        if not os.path.exists(file_path):
            return {"error": "文件不存在", "success": False}
        
        file_size = os.path.getsize(file_path)
        if file_size > self.max_file_size:
            return {"error": f"文件过大（{file_size/1024/1024:.2f}MB，最大支持{self.max_file_size/1024/1024}MB）", "success": False}
        
        file_ext = os.path.splitext(file_path)[1].lower()
        if file_ext not in self.supported_extensions:
            return {"error": f"不支持的文件类型: {file_ext}", "success": False}
        
        try:
            if file_ext == '.pdf':
                return self._parse_pdf(file_path)
            elif file_ext == '.txt':
                return self._parse_txt(file_path)
            elif file_ext in {'.doc', '.docx'}:
                return self._parse_word(file_path)
            elif file_ext in {'.xls', '.xlsx'}:
                return self._parse_excel(file_path)
            elif file_ext == '.md':
                return self._parse_markdown(file_path)
            else:
                return {"error": f"未实现的文件类型: {file_ext}", "success": False}
        except Exception as e:
            logger.error(f"解析文件失败 {file_path}: {e}")
            return {"error": f"解析失败: {str(e)}", "success": False}
    
    def _parse_pdf(self, file_path: str) -> Dict[str, Any]:
        """解析 PDF 文件"""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            logger.warning("PyMuPDF 未安装，使用模拟解析")
            return self._parse_mock(file_path, "PDF")
        
        try:
            doc = fitz.open(file_path)
            text_content = []
            metadata = {
                "page_count": len(doc),
                "title": doc.metadata.get("title", ""),
                "author": doc.metadata.get("author", ""),
                "subject": doc.metadata.get("subject", ""),
            }
            
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                text_content.append({
                    "page": page_num + 1,
                    "content": text
                })
            
            doc.close()
            
            full_text = "\n\n".join([page["content"] for page in text_content])
            
            return {
                "success": True,
                "content": full_text,
                "pages": text_content,
                "metadata": metadata,
                "file_type": "pdf"
            }
            
        except Exception as e:
            logger.error(f"PDF 解析失败: {e}")
            return {"error": f"PDF 解析失败: {str(e)}", "success": False}
    
    def _parse_txt(self, file_path: str) -> Dict[str, Any]:
        """解析 TXT 文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            return {
                "success": True,
                "content": content,
                "metadata": {
                    "char_count": len(content),
                    "line_count": content.count('\n') + 1
                },
                "file_type": "txt"
            }
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='gbk') as f:
                    content = f.read()
                return {
                    "success": True,
                    "content": content,
                    "metadata": {
                        "char_count": len(content),
                        "line_count": content.count('\n') + 1
                    },
                    "file_type": "txt"
                }
            except Exception as e:
                return {"error": f"TXT 解析失败: {str(e)}", "success": False}
        except Exception as e:
            return {"error": f"TXT 解析失败: {str(e)}", "success": False}
    
    def _parse_word(self, file_path: str) -> Dict[str, Any]:
        """解析 Word 文件"""
        try:
            from docx import Document
        except ImportError:
            logger.warning("python-docx 未安装，使用模拟解析")
            return self._parse_mock(file_path, "Word")
        
        try:
            doc = Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            content = "\n\n".join(paragraphs)
            
            metadata = {
                "paragraph_count": len(paragraphs),
                "author": doc.core_properties.author or "",
                "title": doc.core_properties.title or "",
            }
            
            return {
                "success": True,
                "content": content,
                "paragraphs": paragraphs,
                "metadata": metadata,
                "file_type": "word"
            }
        except Exception as e:
            logger.error(f"Word 解析失败: {e}")
            return {"error": f"Word 解析失败: {str(e)}", "success": False}
    
    def _parse_excel(self, file_path: str) -> Dict[str, Any]:
        """解析 Excel 文件"""
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl 未安装，使用模拟解析")
            return self._parse_mock(file_path, "Excel")
        
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            sheets_data = []
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
                sheets_data.append({
                    "sheet_name": sheet_name,
                    "rows": rows,
                    "row_count": len(rows),
                    "col_count": len(rows[0]) if rows else 0
                })
            
            wb.close()
            
            # 生成文本内容
            content_parts = []
            for sheet in sheets_data:
                content_parts.append(f"## {sheet['sheet_name']}")
                content_parts.append("\n".join([" | ".join(row) for row in sheet['rows']]))
            
            content = "\n\n".join(content_parts)
            
            return {
                "success": True,
                "content": content,
                "sheets": sheets_data,
                "metadata": {
                    "sheet_count": len(sheets_data)
                },
                "file_type": "excel"
            }
        except Exception as e:
            logger.error(f"Excel 解析失败: {e}")
            return {"error": f"Excel 解析失败: {str(e)}", "success": False}
    
    def _parse_markdown(self, file_path: str) -> Dict[str, Any]:
        """解析 Markdown 文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # 提取标题
            headers = re.findall(r'^(#{1,6})\s+(.+)$', content, re.MULTILINE)
            
            return {
                "success": True,
                "content": content,
                "metadata": {
                    "header_count": len(headers),
                    "char_count": len(content)
                },
                "file_type": "markdown"
            }
        except Exception as e:
            return {"error": f"Markdown 解析失败: {str(e)}", "success": False}
    
    def _parse_mock(self, file_path: str, file_type: str) -> Dict[str, Any]:
        """模拟解析（当依赖库不可用时）"""
        return {
            "success": True,
            "content": f"这是 {file_type} 文件的模拟解析内容。实际使用时请安装相应的依赖库。",
            "metadata": {
                "file_type": file_type,
                "note": "模拟解析"
            },
            "file_type": file_type.lower()
        }

class TextChunker:
    """文本分块器"""
    
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
    
    def chunk_text(self, text: str) -> List[str]:
        """将文本分块"""
        if not text:
            return []
        
        # 按段落分割
        paragraphs = text.split('\n\n')
        chunks = []
        current_chunk = ""
        
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            
            # 如果当前块加上新段落超过大小限制，保存当前块
            if len(current_chunk) + len(para) > self.chunk_size and current_chunk:
                chunks.append(current_chunk.strip())
                # 保留重叠部分
                current_chunk = current_chunk[-self.chunk_overlap:] if self.chunk_overlap > 0 else ""
            
            current_chunk += para + "\n\n"
        
        # 添加最后一个块
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
        
        # 如果没有分块，直接返回整个文本
        if not chunks and text:
            chunks = [text]
        
        return chunks
    
    def chunk_by_sentences(self, text: str) -> List[str]:
        """按句子分块"""
        sentences = re.split(r'[.!?。！？\n]+', text)
        sentences = [s.strip() for s in sentences if s.strip()]
        
        chunks = []
        current_chunk = ""
        
        for sentence in sentences:
            if len(current_chunk) + len(sentence) > self.chunk_size and current_chunk:
                chunks.append(current_chunk.strip())
                current_chunk = current_chunk[-self.chunk_overlap:] if self.chunk_overlap > 0 else ""
            
            current_chunk += sentence + "。"
        
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
        
        return chunks if chunks else [text]

class TextCleaner:
    """文本清洗器"""
    
    def __init__(self):
        self.patterns = {
            'whitespace': r'\s+',
            'special_chars': r'[^\w\s\u4e00-\u9fff.,!?;:()\'"【】（）。，！？；：]',
            'page_numbers': r'\n\s*第\s*\d+\s*页\s*\n',
            'headers_footers': r'\n\s*(第\s*\d+\s*页|Page\s*\d+|共\s*\d+\s*页)\s*\n',
        }
    
    def clean(self, text: str, aggressive: bool = False) -> str:
        """清洗文本"""
        if not text:
            return ""
        
        cleaned = text
        
        # 移除页码
        cleaned = re.sub(self.patterns['page_numbers'], '\n', cleaned)
        
        # 移除页眉页脚
        cleaned = re.sub(self.patterns['headers_footers'], '\n', cleaned)
        
        # 标准化空白字符
        cleaned = re.sub(self.patterns['whitespace'], ' ', cleaned)
        
        # 移除多余空行
        cleaned = re.sub(r'\n\s*\n\s*\n', '\n\n', cleaned)
        
        if aggressive:
            # 激进模式：移除特殊字符
            cleaned = re.sub(self.patterns['special_chars'], '', cleaned)
        
        return cleaned.strip()
    
    def remove_empty_lines(self, text: str) -> str:
        """移除空行"""
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        return '\n'.join(lines)

# 全局实例
document_parser = DocumentParser()
text_chunker = TextChunker()
text_cleaner = TextCleaner()